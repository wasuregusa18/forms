from django.http.request import QueryDict
from django.http.response import Http404
from django.shortcuts import redirect, render
from django.http import HttpResponse, HttpResponse
from django.views.generic import View
from django.conf import settings
from django.contrib.auth.mixins import LoginRequiredMixin

import re
from urllib.parse import quote

from .. import models, forms


TEMPLATE_ERROR_MESSAGE = "There was a problem uploading your template. Please check your information and try again."

class CreateDocument(LoginRequiredMixin, View):
    model = models.Document
    form_class = forms.DocumentForm
    template_name = "forms/document_form.html"
    context_dict = {"props": sorted(models.CUSTOMER_CHOOSABLE_PROPERTIES)}

    def get(self, request, *args, **kwargs):
        form = self.form_class
        return render(request, self.template_name, self.context_dict)

    def extract_mapping(self, q: QueryDict) -> str:
        """extracts mappings from form's post request
        mapping schema is  "('Cell',customer.attr);('Cell2',customer.attr2)" """
        is_mapping = lambda x: re.search("^(?:cell|val)\d+$", x)
        mapping_keys = filter(is_mapping, q.keys())
        assert len(list(mapping_keys))%2 ==0, "Uneven number of cells and values"

        mapping = []
        for cell_key in mapping_keys:
            val_key = next(mapping_keys)  # iterating 2 at a time
            assert self.is_valid_mapping(cell_key,val_key,q),"Mapping is not valid"
            cell, val = q[cell_key], q[val_key]
            entry = f"('{cell}',customer.{val})"
            mapping.append(entry)
        mapping = ";".join(mapping)
        return mapping

    def is_valid_mapping(self,cell_key:str, val_key:str, q:QueryDict) -> bool:
        cell_match = re.search("^cell(?P<num>\d+)$",cell_key) 
        val_match = re.search(f"^val{cell_match['num']}$",val_key)
        if cell_match and val_match:
            cell, val = q[cell_key], q[val_key]
            is_cell_valid = re.search("^[A-Z]+\d+$",cell)
            is_val_valid = val in models.CUSTOMER_CHOOSABLE_PROPERTIES
            if is_cell_valid and is_val_valid: return True
        return False

    def post(self, request, *args, **kwargs):
        #this needs to be rethought
    
        try:
            mapping = self.extract_mapping(request.POST)
        except AssertionError as e:
            return render(request, self.template_name, self.context_dict) 

        try:
            new_document = models.Document(template=request.FILES["template"], mapping=mapping)
            new_document.save()
            return redirect("forms:index")

        except:
            # new_context_dict = {'error': TEMPLATE_ERROR_MESSAGE}.update(self.context_dict)
            return render(request, self.template_name, self.context_dict)  # this is bugging out


from openpyxl import load_workbook
from zipfile import ZipFile
class DownloadDocument(LoginRequiredMixin, View):
    template_name = "forms/download_document.html"
    form_class = forms.GenerateDocumentForm

    def get(self, request, *args, **kwargs):
        form = self.form_class
        return render(request, self.template_name, {"form": form})

    def post(self, request, *args, **kwargs):
        form = self.form_class(request.POST)
        if form.is_valid():
            print(request.POST)
            document = models.Document.objects.get(pk=request.POST.get("document"))
            customers = [models.Customer.objects.get(pk=pk) for pk in request.POST.getlist("customer")]
            
            paths = []
            for customer in customers:
                paths.append(document.generate_document(customer))
            return self.download(paths)

        return render(request, self.template_name, {"form": form})

    def add_content_disposition_header(self, response, filename):
        """
        Add an RFC5987 / RFC6266 compliant Content-Disposition header to an
        HttpResponse to tell the browser to save the HTTP response to a file.
        """
        try:
            filename.encode('ascii')
            file_expr = 'filename="{}"'.format(filename)
        except UnicodeEncodeError:
            file_expr = "filename*=utf-8''{}".format(quote(filename))
        response['Content-Disposition'] = 'attachment; {}'.format(file_expr)
        return response

    
    def download(self,paths:list)-> HttpResponse:
        """
        return repsonse with zip of files
        """
        file_paths = map(lambda path: settings.MEDIA_ROOT / path, paths)
        checked_file_paths = list(filter(lambda path: path.exists(),file_paths))
        if len(checked_file_paths) == len(paths):
            zip_file_path = settings.MEDIA_ROOT / "generated/customer_data.zip"
            #build zip file
            with ZipFile(zip_file_path,'w') as zip: 
                for file in checked_file_paths: 
                    zip.write(file, file.name) 
                    file.unlink()
            #send reponse with zip file attached
            with open(zip_file_path, "rb") as to_download:
                response = HttpResponse(to_download.read(), content_type= "application/zip")
                response = self.add_content_disposition_header(response,zip_file_path.name)
                zip_file_path.unlink()
                return response
        raise Http404