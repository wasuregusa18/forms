from django.db import models
from django.utils import timezone
from django.urls import reverse
from django.conf import settings
from phonenumber_field.modelfields import PhoneNumberField


import os
from openpyxl import load_workbook
from . import helper_functions


class Customer(models.Model):
    """ model for customer - contains standard fields 
    i.e. name, email_address, phone number, address """
    comp_name = models.CharField(max_length=50)
    email_address = models.EmailField(max_length=300, unique=True)
    phone_number = PhoneNumberField()
    email_hash = models.CharField(max_length=300, primary_key=True)

    # Address
    street1 = models.CharField(max_length=300)
    street2 = models.CharField(max_length=300, blank=True)
    city = models.CharField(max_length=50)
    state = models.CharField(max_length=50)
    post_code = models.CharField(max_length=15)
    country = models.CharField(max_length=50)
    
    @property
    def address(self): 
        return f"""{self.street1}, {self.street2+", " if self.street2 else ""}{self.city},{self.state},{self.country.upper()},{self.post_code}"""

    @property
    def date(self): return timezone.now().strftime("%Y-%m-%d")

    # Internal Information
    entry_date = models.DateField(default=timezone.now)
    last_modified = models.DateField(auto_now=True)

    def __str__(self):
        return f"{self.comp_name}"

    def get_absolute_url(self):
        return reverse("forms:thank_you")

    def save(self, *args, **kwargs):
        if not self.email_hash:
            self.email_hash = helper_functions.hash(self.email_address)
        super(Customer, self).save(*args, **kwargs)

    
    #Japanese specific
    @property
    def jap_comp_name(self): return helper_functions.to_jap_com_name(self.comp_name)
    
    @property
    def jap_address(self):
        street_num, remainder = helper_functions.extract_street_num(self.street1)
        jap_rem = helper_functions.english_to_katakana(remainder)
        jap_city = helper_functions.english_to_katakana(self.city)
        jap_state = helper_functions.english_to_katakana(self.state)
        return f"{self.post_code} {jap_state}　{jap_city}　{jap_rem} {street_num}"

    @property
    def jap_country(self): return helper_functions.to_jap_country(self.country)

    @property
    def jap_date(self):
        now = timezone.now()
        return f"令和{now.year-2018}年{now.month}月{now.day}日"


CUSTOMER_USER_PROPERTIES = set(map(lambda x: x.name, Customer._meta.fields))
CUSTOMER_INTERNAL_PROPERTIES = {"entry_date","last_modified"}
CUSTOMER_GENERATED_PROPERTIES = {
    "jap_comp_name",
    "address",
    "jap_address",
    "date",
    "jap_date",
    "jap_country",
}
CUSTOMER_TOTAL_PROPERTIES = CUSTOMER_USER_PROPERTIES | CUSTOMER_GENERATED_PROPERTIES
CUSTOMER_CHOOSABLE_PROPERTIES =  CUSTOMER_TOTAL_PROPERTIES - CUSTOMER_INTERNAL_PROPERTIES #options to fill in form


class Document(models.Model):
    """ 
    Template to fill in with user data. 
    Mapping between customer fields and template fields 
    excel: "('Cell',customer.attr);('Cell2',customer.attr2)" 
    """
    template = models.FileField(upload_to="temps/", blank=False, null=False)
    mapping = models.CharField(max_length=1000)
    last_modified = models.DateField(auto_now=True)
    
    def __str__(self):
        return os.path.basename(self.template.name)

    def generate_document(self, customer) -> str:
        """takes in customer, builds document based on
        customer data; returns filename of created spreadsheet"""
        workbook = load_workbook(self.template)  
        sheet = workbook.active
        for map in self.mapping.split(";"):
            cell, value = eval(map)
            sheet[cell] = str(value)
        filename = f"{customer.comp_name}-{str(self)}"
        rel_path = "generated/" + filename
        workbook.save(filename=settings.MEDIA_ROOT / rel_path)
        workbook.close()
        return rel_path

